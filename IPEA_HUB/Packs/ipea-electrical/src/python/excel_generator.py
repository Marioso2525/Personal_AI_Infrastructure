"""
excel_generator.py — IPEA_HUB / ipea-electrical
Dashboard Excel de resultados eléctricos: caída de tensión, cédula de cargas,
validación NOM-001-SEDE-2012. Genera un .xlsx multi-hoja profesional.
Requiere: openpyxl
"""

from __future__ import annotations

import json
import math
from dataclasses import dataclass, field
from datetime import datetime
from pathlib import Path
from typing import Any

try:
    from openpyxl import Workbook
    from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
    from openpyxl.chart import BarChart, Reference
    from openpyxl.chart.series import SeriesLabel
    OPENPYXL_AVAILABLE = True
except ImportError:
    OPENPYXL_AVAILABLE = False

COLOR_NAVY   = "1A1A6C"
COLOR_BLUE   = "2E4099"
COLOR_LIGHT  = "C5CAE9"
COLOR_GREEN  = "D4EDDA"
COLOR_RED    = "F8D7DA"
COLOR_YELLOW = "FFF3CD"
COLOR_WHITE  = "FFFFFF"
COLOR_GRAY   = "F5F5F5"


def _f(bold=False, size=10, color="000000", italic=False) -> Font:
    return Font(bold=bold, size=size, color=color, name="Calibri", italic=italic)

def _fill(color: str) -> PatternFill:
    return PatternFill("solid", fgColor=color)

def _align(h="center", v="center", wrap=True) -> Alignment:
    return Alignment(horizontal=h, vertical=v, wrap_text=wrap)

def _border(style="thin", color="9FA8DA") -> Border:
    s = Side(style=style, color=color)
    return Border(left=s, right=s, top=s, bottom=s)

def _header_row(ws, row: int, cols_values: list[tuple[str, str]],
                bg: str = COLOR_BLUE, fg: str = COLOR_WHITE) -> None:
    for col, val in cols_values:
        c = ws[f"{col}{row}"]
        c.value = val
        c.fill = _fill(bg)
        c.font = _f(bold=True, color=fg, size=9)
        c.alignment = _align()
        c.border = _border()

def _data_row(ws, row: int, cols_values: list[tuple[str, Any]],
              bg: str = COLOR_WHITE, align_first_left: bool = True) -> None:
    for i, (col, val) in enumerate(cols_values):
        c = ws[f"{col}{row}"]
        if isinstance(val, float):
            c.value = round(val, 3)
        else:
            c.value = val
        c.fill = _fill(bg)
        c.font = _f(size=9)
        c.alignment = _align(h="left" if (i == 0 and align_first_left) else "center")
        c.border = _border()

def _merge_title(ws, cell_range: str, value: str,
                 bg: str = COLOR_NAVY, fg: str = COLOR_WHITE,
                 size: int = 12, height: int = 22) -> None:
    from openpyxl.utils import range_boundaries
    min_col, min_row, max_col, max_row = range_boundaries(cell_range)
    ws.merge_cells(cell_range)
    ws.row_dimensions[min_row].height = height
    c = ws.cell(row=min_row, column=min_col)
    c.value = value
    c.fill = _fill(bg)
    c.font = _f(bold=True, size=size, color=fg)
    c.alignment = _align()
    c.border = _border()


# ── HOJA 1: Portada / Resumen ───────────────────────────────────────────────
def _build_cover(wb: Workbook, project: dict) -> None:
    ws = wb.active
    ws.title = "Portada"

    for col, w in [("A", 5), ("B", 25), ("C", 20), ("D", 20),
                   ("E", 20), ("F", 20), ("G", 5)]:
        ws.column_dimensions[col].width = w

    _merge_title(ws, "A1:G1", "IPEA_HUB — REPORTE ELÉCTRICO", COLOR_NAVY, size=16, height=35)
    _merge_title(ws, "A2:G2", f"Proyecto: {project.get('name','')}", COLOR_BLUE, size=12, height=20)

    row = 4
    info = [
        ("Número de proyecto:", project.get("number", "")),
        ("Ingeniero responsable:", project.get("engineer", "")),
        ("Fecha de generación:", datetime.now().strftime("%d/%m/%Y %H:%M")),
        ("Norma aplicable:", "NOM-001-SEDE-2012"),
        ("Sistema eléctrico:", project.get("system", "220/127V 3Φ 4H 60Hz")),
    ]
    for label, val in info:
        ws.row_dimensions[row].height = 16
        ws.merge_cells(f"B{row}:C{row}")
        c1 = ws[f"B{row}"]
        c1.value = label
        c1.fill = _fill(COLOR_LIGHT)
        c1.font = _f(bold=True, color=COLOR_NAVY)
        c1.alignment = _align(h="left")
        c1.border = _border()

        ws.merge_cells(f"D{row}:F{row}")
        c2 = ws[f"D{row}"]
        c2.value = val
        c2.alignment = _align(h="left")
        c2.border = _border()
        row += 1

    row += 1
    _merge_title(ws, f"A{row}:G{row}", "CONTENIDO DEL REPORTE", COLOR_BLUE, size=10, height=18)
    row += 1

    sheets_info = [
        ("Hoja 2", "Caída de Tensión", "Resultados por circuito con validación NOM"),
        ("Hoja 3", "Cédula de Cargas", "Listado de cargas por tablero"),
        ("Hoja 4", "Validación NOM", "Verificación integral de circuitos"),
        ("Hoja 5", "Resumen Ejecutivo", "Dashboard con indicadores clave"),
    ]
    _header_row(ws, row, [("B","Hoja"),("C","Nombre"),("D","Descripción")],
                bg=COLOR_BLUE)
    row += 1
    for sheet_num, name, desc in sheets_info:
        bg = COLOR_GRAY if row % 2 == 0 else COLOR_WHITE
        _data_row(ws, row, [("B", sheet_num), ("C", name), ("D", desc)], bg=bg)
        row += 1


# ── HOJA 2: Caída de Tensión ────────────────────────────────────────────────
def _build_voltage_drop_sheet(wb: Workbook, vd_data: list[dict]) -> None:
    ws = wb.create_sheet("Caída de Tensión")

    for col, w in [("A", 28), ("B", 8), ("C", 10), ("D", 10),
                   ("E", 10), ("F", 8), ("G", 8), ("H", 10),
                   ("I", 12), ("J", 10), ("K", 8)]:
        ws.column_dimensions[col].width = w

    _merge_title(ws, "A1:K1", "CÁLCULO DE CAÍDA DE TENSIÓN — NOM-001-SEDE-2012", height=25)
    _merge_title(ws, "A2:K2",
                 "Límites NOM: ≤3% circuito ramal  |  ≤5% total alimentador + ramal",
                 COLOR_BLUE, size=9, height=16)

    row = 4
    headers = [
        ("A","CIRCUITO / DESCRIPCIÓN"), ("B","VOLTAJE\n(V)"),
        ("C","CORRIENTE\n(A)"), ("D","LONGITUD\n(m)"),
        ("E","CALIBRE"), ("F","FP"), ("G","TIPO"),
        ("H","VD (V)"), ("I","VD (%)"), ("J","LÍMITE\nNOM (%)"),
        ("K","ESTADO"),
    ]
    ws.row_dimensions[row].height = 30
    _header_row(ws, row, headers)
    row += 1

    for i, d in enumerate(vd_data):
        bg = COLOR_GRAY if i % 2 == 0 else COLOR_WHITE
        vd_pct = d.get("vd_percent", 0)
        status = "✅ OK" if vd_pct <= 3.0 else ("⚠️ >3%" if vd_pct <= 5.0 else "❌ >5%")
        status_bg = COLOR_GREEN if vd_pct <= 3.0 else (COLOR_YELLOW if vd_pct <= 5.0 else COLOR_RED)

        ws.row_dimensions[row].height = 14
        cols = [
            ("A", d.get("description","")), ("B", d.get("voltage_V",0)),
            ("C", d.get("current_A",0)), ("D", d.get("length_m",0)),
            ("E", d.get("conductor","")), ("F", d.get("pf",0.85)),
            ("G", d.get("circuit_type","3Φ")), ("H", round(d.get("vd_V",0),3)),
            ("I", round(vd_pct,2)), ("J", 3.0),
        ]
        _data_row(ws, row, cols, bg=bg)

        c = ws[f"K{row}"]
        c.value = status
        c.fill = _fill(status_bg)
        c.font = _f(bold=True, size=9)
        c.alignment = _align()
        c.border = _border()
        row += 1


# ── HOJA 3: Cédula de Cargas ────────────────────────────────────────────────
def _build_load_schedule_sheet(wb: Workbook, loads_data: list[dict]) -> None:
    ws = wb.create_sheet("Cédula de Cargas")

    for col, w in [("A", 30), ("B", 10), ("C", 12), ("D", 8),
                   ("E", 10), ("F", 10), ("G", 10), ("H", 12)]:
        ws.column_dimensions[col].width = w

    _merge_title(ws, "A1:H1", "CÉDULA DE CARGAS — NOM-001-SEDE-2012", height=25)

    row = 3
    headers = [
        ("A","DESCRIPCIÓN DE CARGA"), ("B","TIPO"),
        ("C","CANTIDAD"), ("D","W/UNIDAD"),
        ("E","W TOTAL"), ("F","F.P."),
        ("G","CORRIENTE\n(A)"), ("H","¿CONTINUA?"),
    ]
    ws.row_dimensions[row].height = 28
    _header_row(ws, row, headers)
    row += 1

    total_w = 0.0
    for i, load in enumerate(loads_data):
        bg = COLOR_GRAY if i % 2 == 0 else COLOR_WHITE
        total_w_load = load.get("quantity", 0) * load.get("unit_W", 0)
        total_w += total_w_load
        curr = total_w_load / (load.get("voltage_V", 127) * load.get("pf", 0.9))

        ws.row_dimensions[row].height = 14
        _data_row(ws, row, [
            ("A", load.get("name","")),
            ("B", load.get("type","")),
            ("C", load.get("quantity",0)),
            ("D", load.get("unit_W",0)),
            ("E", round(total_w_load,1)),
            ("F", load.get("pf",0.9)),
            ("G", round(curr,2)),
            ("H", "SÍ" if load.get("continuous") else "NO"),
        ], bg=bg)
        row += 1

    ws.row_dimensions[row].height = 16
    _data_row(ws, row, [
        ("A","TOTAL"), ("B",""), ("C",""), ("D",""),
        ("E", round(total_w,1)), ("F",""), ("G",""), ("H",""),
    ], bg=COLOR_LIGHT)
    for col in "ABCDEFGH":
        ws[f"{col}{row}"].font = _f(bold=True, color=COLOR_NAVY)


# ── HOJA 4: Validación NOM ──────────────────────────────────────────────────
def _build_validation_sheet(wb: Workbook, validations: list[dict]) -> None:
    ws = wb.create_sheet("Validación NOM")

    for col, w in [("A", 28), ("B", 15), ("C", 15), ("D", 12), ("E", 10)]:
        ws.column_dimensions[col].width = w

    _merge_title(ws, "A1:E1", "VALIDACIÓN INTEGRAL NOM-001-SEDE-2012", height=25)

    row = 3
    _header_row(ws, row, [
        ("A","CIRCUITO"), ("B","VERIFICACIÓN"),
        ("C","VALOR"), ("D","LÍMITE"), ("E","ESTADO"),
    ])
    row += 1

    for v in validations:
        circuit_name = v.get("circuit","")
        for item in v.get("items", []):
            bg = COLOR_GRAY if row % 2 == 0 else COLOR_WHITE
            passed = item.get("passed", False)
            status = "✅ PASA" if passed else "❌ FALLA"
            status_bg = COLOR_GREEN if passed else COLOR_RED

            ws.row_dimensions[row].height = 14
            _data_row(ws, row, [
                ("A", circuit_name),
                ("B", item.get("name","")),
                ("C", item.get("value","")),
                ("D", item.get("limit","")),
            ], bg=bg)

            c = ws[f"E{row}"]
            c.value = status
            c.fill = _fill(status_bg)
            c.font = _f(bold=True, size=9)
            c.alignment = _align()
            c.border = _border()
            row += 1


# ── HOJA 5: Resumen Ejecutivo ───────────────────────────────────────────────
def _build_summary_sheet(wb: Workbook, summary: dict) -> None:
    ws = wb.create_sheet("Resumen Ejecutivo")

    for col, w in [("A", 5), ("B", 30), ("C", 20), ("D", 20), ("E", 5)]:
        ws.column_dimensions[col].width = w

    _merge_title(ws, "A1:E1", "RESUMEN EJECUTIVO — INDICADORES CLAVE", height=30)

    row = 3
    _merge_title(ws, f"A{row}:E{row}", "INDICADORES DE DISEÑO", COLOR_BLUE, size=10, height=18)
    row += 1

    kpis = [
        ("Circuitos totales", summary.get("total_circuits", 0), "und"),
        ("Circuitos que cumplen NOM", summary.get("circuits_ok", 0), "und"),
        ("% cumplimiento NOM", summary.get("compliance_pct", 0), "%"),
        ("Carga total instalada", summary.get("total_load_kW", 0), "kW"),
        ("Carga de demanda", summary.get("demand_kW", 0), "kW"),
        ("Factor de demanda", summary.get("demand_factor", 0), ""),
        ("VD máximo encontrado", summary.get("max_vd_pct", 0), "%"),
        ("VD promedio", summary.get("avg_vd_pct", 0), "%"),
    ]

    for label, value, unit in kpis:
        ws.row_dimensions[row].height = 16
        ws.merge_cells(f"B{row}:C{row}")
        c1 = ws[f"B{row}"]
        c1.value = label
        c1.fill = _fill(COLOR_LIGHT)
        c1.font = _f(bold=True, color=COLOR_NAVY)
        c1.alignment = _align(h="left")
        c1.border = _border()

        ws[f"D{row}"].value = f"{value} {unit}".strip()
        ws[f"D{row}"].alignment = _align()
        ws[f"D{row}"].border = _border()
        row += 1

    row += 1
    compliance = summary.get("compliance_pct", 0)
    status_text = "✅ PROYECTO CUMPLE NOM-001-SEDE-2012" if compliance == 100 else f"⚠️ {100-compliance:.0f}% DE CIRCUITOS REQUIEREN CORRECCIÓN"
    status_color = COLOR_GREEN if compliance == 100 else COLOR_RED
    _merge_title(ws, f"A{row}:E{row}", status_text, status_color,
                 fg="000000" if compliance == 100 else COLOR_WHITE, height=25)


# ── FUNCIÓN PRINCIPAL ────────────────────────────────────────────────────────
def generate_electrical_dashboard(
    project: dict,
    vd_results: list[dict],
    loads: list[dict],
    validations: list[dict],
    output_path: str,
) -> str:
    if not OPENPYXL_AVAILABLE:
        return "ERROR: openpyxl no instalado. Ejecutar: pip install openpyxl"

    wb = Workbook()

    _build_cover(wb, project)
    _build_voltage_drop_sheet(wb, vd_results)
    _build_load_schedule_sheet(wb, loads)
    _build_validation_sheet(wb, validations)

    circuits_ok = sum(1 for v in validations if all(i.get("passed") for i in v.get("items",[])))
    total_circuits = len(validations)
    vd_pcts = [d.get("vd_percent", 0) for d in vd_results]

    summary = {
        "total_circuits": total_circuits,
        "circuits_ok": circuits_ok,
        "compliance_pct": round(circuits_ok / total_circuits * 100, 1) if total_circuits else 0,
        "total_load_kW": round(sum(l.get("quantity",0)*l.get("unit_W",0) for l in loads)/1000, 2),
        "demand_kW": round(sum(l.get("quantity",0)*l.get("unit_W",0) for l in loads)*0.75/1000, 2),
        "demand_factor": 0.75,
        "max_vd_pct": round(max(vd_pcts), 2) if vd_pcts else 0,
        "avg_vd_pct": round(sum(vd_pcts)/len(vd_pcts), 2) if vd_pcts else 0,
    }
    _build_summary_sheet(wb, summary)

    out = Path(output_path)
    out.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(out))
    return f"Dashboard eléctrico generado: {out}"


if __name__ == "__main__":
    project = {
        "name": "Edificio Corporativo Norte",
        "number": "IPEA-2025-001",
        "engineer": "Ing. IPEA_HUB",
        "system": "220/127V 3Φ 4H 60Hz",
    }

    vd_results = [
        {"description":"Alimentador T-PB-01","voltage_V":220,"current_A":60,"length_m":45,
         "conductor":"4AWG","pf":0.85,"circuit_type":"3Φ","vd_V":4.6,"vd_percent":2.1},
        {"description":"Circuito Iluminación C-1","voltage_V":127,"current_A":15,"length_m":30,
         "conductor":"12AWG","pf":0.95,"circuit_type":"1Φ","vd_V":3.6,"vd_percent":2.8},
        {"description":"Circuito Contactos C-3","voltage_V":127,"current_A":16,"length_m":50,
         "conductor":"12AWG","pf":0.90,"circuit_type":"1Φ","vd_V":6.1,"vd_percent":4.8},
    ]

    loads = [
        {"name":"Iluminación LED","type":"lighting","quantity":40,"unit_W":60,"pf":0.95,"voltage_V":127},
        {"name":"Contactos Oficina","type":"receptacle","quantity":30,"unit_W":180,"pf":0.90,"voltage_V":127},
        {"name":"Aire Acondicionado 5TR","type":"hvac","quantity":3,"unit_W":6000,"pf":0.85,"voltage_V":220},
    ]

    validations = [
        {"circuit":"Alimentador T-PB-01","items":[
            {"name":"Ampacidad","value":"85A","limit":"≥75A","passed":True},
            {"name":"Caída Tensión","value":"2.1%","limit":"≤3%","passed":True},
            {"name":"Protección","value":"70A","limit":"≤85A","passed":True},
        ]},
        {"circuit":"Circuito C-3","items":[
            {"name":"Ampacidad","value":"25A","limit":"≥20A","passed":True},
            {"name":"Caída Tensión","value":"4.8%","limit":"≤3%","passed":False},
            {"name":"Protección","value":"20A","limit":"≤25A","passed":True},
        ]},
    ]

    msg = generate_electrical_dashboard(project, vd_results, loads, validations,
                                        "/tmp/dashboard_electrico.xlsx")
    print(msg)
